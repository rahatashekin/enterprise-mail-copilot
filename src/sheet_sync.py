"""
Business Mail Copilot — Live Tracker & Sheet Synchronizer
Enterprise Trading Corporation

Production Dual-Sheet Live Tracking Engine (Multi-Tier):
  1. Sheet 1: 'Project Master Pipeline' (Upsert Mode)
     - Exactly one row per unique Project ID.
     - Supports 3 tiers: Buyer (Tender lifecycle), Supplier (Quotation lifecycle), Bank (LC lifecycle).
     - Incremental state accumulation (preserves historical values when later events update a project).
     - Color-coded stage badges, freeze panes, auto-filter, and clickable Gmail deep-links.
  2. Sheet 2: 'Tender Event Log' (Append-Only Audit Trail)
     - Chronological record of every email interaction across all projects and tiers.
     - Includes timestamp, stage, summary, sender, tier, critical action, and direct email hyperlink.
"""

import os
import time
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, Optional, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"
EXCEL_PATH = DATA_DIR / "Tender_Tracker.xlsx"

# ------------------------------------------------------------------ #
# Column Schemas                                                     #
# ------------------------------------------------------------------ #

MASTER_COLUMNS = [
    "Sl", "Project ID", "Project Title", "Source / Entity",
    "Status", "Reference No.", "Amount (BDT)",
    "Security / Deposit (BDT)", "Key Date / Deadline", "Item Summary",
    "Latest Critical Action", "Contact / Source", "Last Updated", "Gmail Link",
    "Tier", "Domain"
]

EVENT_COLUMNS = [
    "Sl", "Event Timestamp", "Project ID", "Source / Entity",
    "Status", "Subject / Event Summary", "Critical Action",
    "Sender", "Direct Gmail Link", "Tier"
]

# ------------------------------------------------------------------ #
# Visual Styling Standards                                           #
# ------------------------------------------------------------------ #

STAGE_COLORS = {
    # Tier 1: Buyer Tender Lifecycle
    "1_NEW_RFQ": {"fill": "D9E1F2", "font": "1F4E78"},          # Soft Sky Blue / Navy
    "2_NEGOTIATION": {"fill": "FCE4D6", "font": "C65911"},      # Soft Amber / Dark Orange
    "3_PO_AWARD": {"fill": "E2EFDA", "font": "276A3C"},         # Soft Mint Green / Forest Green
    "4_DELIVERY_CHALLAN": {"fill": "EDEDF5", "font": "4F407A"}, # Soft Lavender / Deep Purple
    "5_PG_PAYMENT": {"fill": "FFF2CC", "font": "806000"},       # Soft Golden Yellow / Bronze
    # Tier 2: Supplier Quotation Lifecycle
    "S1_QUOTE_RECEIVED": {"fill": "E2F0D9", "font": "375623"},  # Light Green / Dark Green
    "S2_UNDER_REVIEW": {"fill": "DEEBF7", "font": "2E5B9A"},    # Light Blue / Mid Blue
    "S3_ORDER_PLACED": {"fill": "C6EFCE", "font": "006100"},    # Green / Dark Green
    "S4_AWAITING_DELIVERY": {"fill": "FCE4D6", "font": "974706"},# Peach / Brown
    # Tier 3: Bank/Finance Lifecycle
    "B1_LC_DRAFT": {"fill": "F2DCDB", "font": "953735"},        # Light Rose / Dark Red
    "B2_LC_OPENED": {"fill": "DAEEF3", "font": "205867"},       # Light Teal / Dark Teal
    "B3_DOCS_SUBMITTED": {"fill": "E4DFEC", "font": "604A7B"},  # Light Purple / Purple
    "B4_PAYMENT_RECEIVED": {"fill": "D8E4BC", "font": "4F6228"},# Light Olive / Olive
}

DARK_NAVY_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
CENTER_ALIGN_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_ALIGN_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT_ALIGN_VAL = Alignment(horizontal="right", vertical="center")
REGULAR_FONT = Font(name="Calibri", size=10)
BOLD_FONT = Font(name="Calibri", size=10, bold=True)
LINK_FONT = Font(name="Calibri", size=10, color="0563C1", underline="single")

HEADER_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='1F4E78'),
    bottom=Side(style='medium', color='1F4E78')
)

CELL_BORDER = Border(
    left=Side(style='thin', color='E0E0E0'),
    right=Side(style='thin', color='E0E0E0'),
    top=Side(style='thin', color='E0E0E0'),
    bottom=Side(style='thin', color='E0E0E0')
)


class SheetSync:
    """Synchronizes tender lifecycle events into a professional dual-sheet Excel tracker."""

    def __init__(self, excel_path: Path = EXCEL_PATH):
        self.excel_path = Path(excel_path)
        self._ensure_workbook()

    def _ensure_workbook(self):
        """Creates or upgrades the tracker workbook to dual-sheet architecture."""
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        if not self.excel_path.exists():
            wb = openpyxl.Workbook()
            # Rename default sheet to Master Pipeline
            ws_master = wb.active
            ws_master.title = "Project Master Pipeline"
            self._init_sheet_headers(ws_master, MASTER_COLUMNS)

            # Create second sheet for Event Log
            ws_events = wb.create_sheet(title="Tender Event Log")
            self._init_sheet_headers(ws_events, EVENT_COLUMNS)

            wb.save(self.excel_path)
        else:
            try:
                wb = openpyxl.load_workbook(self.excel_path)
                modified = False

                if "Project Master Pipeline" not in wb.sheetnames:
                    ws_master = wb.create_sheet(title="Project Master Pipeline", index=0)
                    self._init_sheet_headers(ws_master, MASTER_COLUMNS)
                    modified = True

                if "Tender Event Log" not in wb.sheetnames:
                    ws_events = wb.create_sheet(title="Tender Event Log", index=1)
                    self._init_sheet_headers(ws_events, EVENT_COLUMNS)
                    modified = True

                # Remove legacy single-sheet if present and new sheets are active
                if "Tender Radar" in wb.sheetnames and "Project Master Pipeline" in wb.sheetnames:
                    del wb["Tender Radar"]
                    modified = True

                if modified:
                    wb.save(self.excel_path)
            except Exception:
                # If corrupted or unreadable, create a clean workbook
                wb = openpyxl.Workbook()
                ws_master = wb.active
                ws_master.title = "Project Master Pipeline"
                self._init_sheet_headers(ws_master, MASTER_COLUMNS)
                ws_events = wb.create_sheet(title="Tender Event Log")
                self._init_sheet_headers(ws_events, EVENT_COLUMNS)
                wb.save(self.excel_path)

    def _init_sheet_headers(self, ws, columns: List[str]):
        """Applies uniform professional header styling, freeze panes, and filters."""
        ws.views.sheetView[0].showGridLines = True
        ws.append(columns)
        ws.row_dimensions[1].height = 30

        for col_idx in range(1, len(columns) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = HEADER_FONT
            cell.fill = DARK_NAVY_FILL
            cell.alignment = CENTER_ALIGN_WRAP
            cell.border = HEADER_BORDER

        ws.freeze_panes = "A2"
        # Set auto-filter for all header columns
        col_letter = get_column_letter(len(columns))
        ws.auto_filter.ref = f"A1:{col_letter}1"

    def sync_event(
        self,
        extracted_data: Dict[str, Any],
        project_summary: Optional[Dict[str, Any]] = None,
        message_id: str = "",
        thread_id: str = "",
        tier: str = "BUYER",
        source_domain: str = ""
    ) -> str:
        """
        Atomically synchronizes an email event into:
          1. 'Project Master Pipeline' (Upsert unique row per project)
          2. 'Tender Event Log' (Append-only audit trail)
        Supports multi-tier: BUYER, SUPPLIER, BANK.
        """
        self._ensure_workbook()
        wb = openpyxl.load_workbook(self.excel_path)
        ws_master = wb["Project Master Pipeline"]
        ws_events = wb["Tender Event Log"]

        proj_summary = project_summary or {}
        now_str = datetime.now().strftime("%Y-%m-%d %H:%M")

        # Resolve core identifiers
        project_id = (
            extracted_data.get("project_id") or
            proj_summary.get("project_id") or
            "PROJ-UNKNOWN"
        )
        project_title = (
            proj_summary.get("title") or
            extracted_data.get("project_title") or
            extracted_data.get("tender_title") or
            "N/A"
        )
        source_entity = (
            extracted_data.get("plant_name") or
            extracted_data.get("supplier_name") or
            extracted_data.get("bank_name") or
            proj_summary.get("plant_name") or
            source_domain or
            "অজ্ঞাত উৎস"
        )
        stage = extracted_data.get("lifecycle_stage", "1_NEW_RFQ")

        # Resolve official reference (Tender / PO / Challan No)
        tender_no = (
            extracted_data.get("tender_no") or
            proj_summary.get("tender_no") or
            ""
        )
        po_no = (
            extracted_data.get("po_number") or
            extracted_data.get("po_no") or
            proj_summary.get("po_no") or
            ""
        )
        if po_no and tender_no and tender_no != "N/A":
            official_ref = f"PO: {po_no} (Tender: {tender_no})"
        elif po_no:
            official_ref = f"PO: {po_no}"
        elif tender_no:
            official_ref = tender_no
        else:
            official_ref = extracted_data.get("enquiry_no") or "প্রযোজ্য নয়"

        # Resolve commercial & delivery values
        po_val = (
            extracted_data.get("po_value_bdt") or
            extracted_data.get("total_contract_value") or
            proj_summary.get("contract_value") or
            "N/A"
        )
        emd_val = (
            extracted_data.get("tender_security_bdt") or
            extracted_data.get("tender_security") or
            extracted_data.get("release_amount_bdt") or
            proj_summary.get("pg_amount") or
            "N/A"
        )
        deadline = (
            extracted_data.get("revised_deadline") or
            extracted_data.get("submission_deadline") or
            extracted_data.get("delivery_timeframe") or
            extracted_data.get("delivery_location") or
            "N/A"
        )
        boq_summary = (
            extracted_data.get("items_summary") or
            extracted_data.get("approved_items") or
            extracted_data.get("quantities") or
            "N/A"
        )
        critical_action = extracted_data.get("critical_action", "রিভিউ প্রয়োজন")
        officer = (
            extracted_data.get("issuing_officer") or
            extracted_data.get("officer_name") or
            "N/A"
        )
        msg_id = message_id or extracted_data.get("message_id", "")

        # -------------------------------------------------------------- #
        # 1. UPSERT: Project Master Pipeline                             #
        # -------------------------------------------------------------- #
        existing_row = None
        for r in range(2, ws_master.max_row + 1):
            cell_val = ws_master.cell(row=r, column=2).value
            if cell_val and str(cell_val).strip() == str(project_id).strip():
                existing_row = r
                break

        if existing_row:
            target_row = existing_row
            sl_no = ws_master.cell(row=existing_row, column=1).value or (existing_row - 1)
            # Incremental state accumulation: retain existing values if new event is "N/A"
            old_title = ws_master.cell(row=existing_row, column=3).value
            old_entity = ws_master.cell(row=existing_row, column=4).value
            old_ref = ws_master.cell(row=existing_row, column=6).value
            old_po_val = ws_master.cell(row=existing_row, column=7).value
            old_emd = ws_master.cell(row=existing_row, column=8).value
            old_deadline = ws_master.cell(row=existing_row, column=9).value
            old_boq = ws_master.cell(row=existing_row, column=10).value
            old_officer = ws_master.cell(row=existing_row, column=12).value

            final_title = project_title if (project_title and project_title != "N/A") else old_title
            final_entity = source_entity if (source_entity and source_entity != "অজ্ঞাত উৎস") else old_entity
            final_ref = official_ref if (official_ref and official_ref != "প্রযোজ্য নয়") else old_ref
            final_po_val = po_val if (po_val and po_val != "N/A") else old_po_val
            final_emd = emd_val if (emd_val and emd_val != "N/A") else old_emd
            final_deadline = deadline if (deadline and deadline != "N/A") else old_deadline
            final_boq = boq_summary if (boq_summary and boq_summary != "N/A") else old_boq
            final_officer = officer if (officer and officer != "N/A") else old_officer
        else:
            target_row = ws_master.max_row + 1
            sl_no = target_row - 1
            final_title = project_title
            final_entity = source_entity
            final_ref = official_ref
            final_po_val = po_val
            final_emd = emd_val
            final_deadline = deadline
            final_boq = boq_summary
            final_officer = officer

        master_row_values = [
            sl_no,
            project_id,
            final_title,
            final_entity,
            stage,
            final_ref,
            final_po_val,
            final_emd,
            final_deadline,
            final_boq,
            critical_action,
            final_officer,
            now_str,
            "Open in Gmail",
            tier,
            source_domain
        ]

        ws_master.row_dimensions[target_row].height = 36
        for col_idx, val in enumerate(master_row_values, 1):
            cell = ws_master.cell(row=target_row, column=col_idx)
            cell.value = val
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER

            # Specific column alignments
            if col_idx in [1, 2, 4, 5, 9, 13, 14, 15, 16]:
                cell.alignment = CENTER_ALIGN_WRAP
            elif col_idx in [7, 8]:
                cell.alignment = RIGHT_ALIGN_VAL
            else:
                cell.alignment = LEFT_ALIGN_WRAP

            # Project ID emphasis
            if col_idx == 2:
                cell.font = BOLD_FONT

            # Stage Color Badge on Col 5
            if col_idx == 5:
                cfg = STAGE_COLORS.get(stage, {"fill": "F2F2F2", "font": "000000"})
                cell.fill = PatternFill(start_color=cfg["fill"], end_color=cfg["fill"], fill_type="solid")
                cell.font = Font(name="Calibri", size=10, bold=True, color=cfg["font"])

            # PO Value emphasis
            if col_idx == 7 and val != "N/A":
                cell.font = Font(name="Calibri", size=10, bold=True, color="276A3C")

            # Clickable Gmail Hyperlink on Col 14
            if col_idx == 14 and msg_id:
                cell.font = LINK_FONT
                cell.hyperlink = f"https://mail.google.com/mail/u/0/#inbox/{msg_id}"

        # -------------------------------------------------------------- #
        # 2. APPEND: Tender Event Log                                    #
        # -------------------------------------------------------------- #
        event_row = ws_events.max_row + 1
        event_sl = event_row - 1
        event_date = extracted_data.get("date") or now_str
        subject_summary = extracted_data.get("subject") or final_title
        sender = extracted_data.get("sender") or "Buyer / Plant Authority"

        event_values = [
            event_sl,
            event_date,
            project_id,
            source_entity,
            stage,
            subject_summary,
            critical_action,
            sender,
            "View Email",
            tier
        ]

        ws_events.row_dimensions[event_row].height = 26
        for col_idx, val in enumerate(event_values, 1):
            cell = ws_events.cell(row=event_row, column=col_idx)
            cell.value = val
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER

            if col_idx in [1, 2, 3, 4, 5, 9, 10]:
                cell.alignment = CENTER_ALIGN_WRAP
            else:
                cell.alignment = LEFT_ALIGN_WRAP

            # Stage Color Badge on Col 5
            if col_idx == 5:
                cfg = STAGE_COLORS.get(stage, {"fill": "F2F2F2", "font": "000000"})
                cell.fill = PatternFill(start_color=cfg["fill"], end_color=cfg["fill"], fill_type="solid")
                cell.font = Font(name="Calibri", size=10, bold=True, color=cfg["font"])

            # Clickable Gmail Hyperlink on Col 9
            if col_idx == 9 and msg_id:
                cell.font = LINK_FONT
                cell.hyperlink = f"https://mail.google.com/mail/u/0/#inbox/{msg_id}"

        # -------------------------------------------------------------- #
        # Auto-adjust column widths dynamically                          #
        # -------------------------------------------------------------- #
        self._autofit_columns(ws_master)
        self._autofit_columns(ws_events)

        wb.save(self.excel_path)
        return str(self.excel_path)

    def _autofit_columns(self, ws):
        """Auto-fits column widths with bounded min/max values for readability."""
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for c in col:
                val = str(c.value or '')
                # If cell contains multiple lines, calculate based on the longest line
                for line in val.split('\n'):
                    max_len = max(max_len, len(line))
            ws.column_dimensions[col_letter].width = max(min(max_len + 4, 50), 12)

    def append_tender(self, tender_data: Dict[str, Any]) -> str:
        """Backward-compatible wrapper for single-turn calls."""
        return self.sync_event(
            extracted_data=tender_data,
            project_summary=None,
            message_id=tender_data.get("message_id", "")
        )

